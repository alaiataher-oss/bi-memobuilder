from __future__ import annotations

from pathlib import Path
from typing import Any

from .config import ROOT
from .rules_loader import load_templates

FONTS_DIR = ROOT / "assets" / "fonts"
STATIC_FONTS_DIR = ROOT / "static" / "fonts"

# Prefer BI-licensed Frutiger if present; else open lookalike so preview == download.
BODY_CANDIDATES = (
    "Frutiger45Light.ttf",
    "Frutiger45Light.otf",
    "Frutiger 45 Light.ttf",
    "Frutiger 45 Light.otf",
    "Frutiger.ttf",
    "Frutiger.otf",
    "SourceSans3-Light.ttf",
    "SourceSans3-Light.otf",
    "SourceSans3-Regular.ttf",
    "SourceSans3-Regular.otf",
)
HEADING_CANDIDATES = (
    "Optima.ttf",
    "Optima.otf",
    "OptimaBold.ttf",
    "SourceSerif4-Bold.ttf",
    "SourceSerif4-Regular.ttf",
)


def _first_existing(names: tuple[str, ...]) -> Path | None:
    for name in names:
        path = FONTS_DIR / name
        if path.exists() and path.is_file() and path.stat().st_size > 1000:
            return path
    return None


def resolve_font_paths() -> dict[str, Path | None]:
    pack = load_templates()
    fonts = pack.get("styles", {}).get("fonts", {})
    configured_body = ROOT / fonts.get("frutiger_asset", "assets/fonts/Frutiger.ttf")
    configured_head = ROOT / fonts.get("optima_asset", "assets/fonts/Optima.ttf")
    body = configured_body if configured_body.exists() else _first_existing(BODY_CANDIDATES)
    heading = configured_head if configured_head.exists() else _first_existing(HEADING_CANDIDATES)
    if heading is None:
        heading = body
    return {"heading": heading, "body": body}


def body_font_family_name(path: Path | None = None) -> str:
    """Word/CSS family name for active body font."""
    p = path or resolve_font_paths()["body"]
    if p is None:
        return "Frutiger 45 Light"
    name = p.name.lower().replace(" ", "")
    if "frutiger" in name:
        return "Frutiger 45 Light"
    if "sourcesans" in name:
        return "Source Sans 3"
    return "Frutiger 45 Light"


def heading_font_family_name(path: Path | None = None) -> str:
    p = path or resolve_font_paths()["heading"]
    if p is None:
        return "Optima"
    name = p.name.lower()
    if "optima" in name:
        return "Optima"
    if "sourceserif" in name:
        return "Source Serif 4"
    if "sourcesans" in name:
        return "Source Sans 3"
    return "Optima"


def font_assets_ready(template_pack: dict[str, Any] | None = None) -> tuple[bool, list[str]]:
    paths = resolve_font_paths()
    missing: list[str] = []
    if paths["body"] is None:
        missing.append(
            "assets/fonts/Frutiger45Light.ttf (lisensi BI) "
            "atau SourceSans3-Light.ttf (stand-in legal)"
        )
    return (paths["body"] is not None, missing)


def sync_static_fonts() -> dict[str, str | None]:
    """Copy active fonts to static/fonts for matching web preview."""
    STATIC_FONTS_DIR.mkdir(parents=True, exist_ok=True)
    paths = resolve_font_paths()
    urls: dict[str, str | None] = {"body": None, "heading": None}
    for key, src in paths.items():
        if not src:
            continue
        dest = STATIC_FONTS_DIR / src.name
        data = src.read_bytes()
        if not dest.exists() or dest.read_bytes() != data:
            dest.write_bytes(data)
        urls[key] = f"/static/fonts/{src.name}"
    return urls


def apply_run_fonts(document, body_name: str, heading_name: str) -> None:
    """Force every run in the DOCX to use the resolved family names."""
    from docx.oxml.ns import qn

    def fix_run(run, name: str) -> None:
        run.font.name = name
        r = run._element
        rPr = r.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), name)
        rFonts.set(qn("w:hAnsi"), name)
        rFonts.set(qn("w:cs"), name)
        rFonts.set(qn("w:eastAsia"), name)

    for p in document.paragraphs:
        style_name = (p.style.name if p.style is not None else "") or ""
        fam = heading_name if style_name in ("MemoTitle",) else body_name
        for run in p.runs:
            fix_run(run, fam if style_name != "MemoTitle" else heading_name)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        fix_run(run, body_name)
