from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY

from .config import ROOT
from .attachments import attachment_is_meaningful, resolve_attachment_path
from .filename import generate_filename
from .fonts import (
    apply_run_fonts,
    body_font_family_name,
    heading_font_family_name,
    resolve_font_paths,
)
from .outline import (
    LEVEL_HANGING_CM,
    LEVEL_INDENT_CM,
    format_attachment_count,
    parse_outline_blocks,
    points_to_content,
    section_heading,
)
from .rules_loader import get_template, load_templates
from .validation import validate_document


def _set_cell_border(cell, **kwargs) -> None:
    """Set cell borders; kwargs keys: top/left/bottom/right -> {sz, color, val}."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        element = tcBorders.find(qn(f"w:{edge}"))
        if element is not None:
            tcBorders.remove(element)
        if edge in kwargs:
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), kwargs[edge].get("val", "single"))
            element.set(qn("w:sz"), str(kwargs[edge].get("sz", 12)))
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), kwargs[edge].get("color", "000000"))
            tcBorders.append(element)


def _clear_cell(cell) -> None:
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)


def _add_centered_run(paragraph, text: str, *, bold=False, underline=False, size_pt=11, font="Frutiger 45 Light"):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.underline = underline
    run.font.size = Pt(size_pt)
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return run


def _fill_accountability_cell(cell, label: str, person: dict[str, str]) -> None:
    """2×2 cell: label header + jabatan + ruang tanda tangan + nama + pangkat."""
    _clear_cell(cell)
    border = {"sz": 12, "color": "000000", "val": "single"}
    _set_cell_border(cell, top=border, left=border, bottom=border, right=border)

    head = cell.add_paragraph()
    _add_centered_run(head, label, bold=True)
    pPr = head._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)

    role = cell.add_paragraph()
    _add_centered_run(role, person.get("title") or "")

    # Generous signature space
    for _ in range(5):
        sp = cell.add_paragraph()
        sp.paragraph_format.space_after = Pt(8)
        _add_centered_run(sp, "")

    name_p = cell.add_paragraph()
    _add_centered_run(name_p, person.get("name") or "", underline=True, bold=True)

    rank_p = cell.add_paragraph()
    _add_centered_run(rank_p, person.get("rank") or "")


def _add_accountability_table(document: Document, labels: list[str], roles: dict[str, Any]) -> None:
    mapping = {
        "Dipersiapkan oleh": "prepared_by",
        "Diperiksa oleh": "reviewed_by",
        "Didukung oleh": "supported_by",
        "Disetujui oleh": "approved_by",
        "Diterima oleh": "received_by",
    }
    # Keep 2×2 grid (two columns); pad odd count with empty cell
    n_rows = (len(labels) + 1) // 2
    table = document.add_table(rows=n_rows, cols=2)
    table.autofit = True
    for i in range(n_rows):
        left_label = labels[i * 2]
        right_label = labels[i * 2 + 1] if i * 2 + 1 < len(labels) else None
        left_person = roles.get(mapping[left_label]) or {}
        _fill_accountability_cell(table.cell(i, 0), left_label, left_person)
        if right_label:
            right_person = roles.get(mapping[right_label]) or {}
            _fill_accountability_cell(table.cell(i, 1), right_label, right_person)
        else:
            _clear_cell(table.cell(i, 1))
            border = {"sz": 12, "color": "000000", "val": "single"}
            _set_cell_border(table.cell(i, 1), top=border, left=border, bottom=border, right=border)


def _accountability_pdf_table(labels: list[str], roles: dict[str, Any], body_font: str, size: float) -> Table:
    mapping = {
        "Dipersiapkan oleh": "prepared_by",
        "Diperiksa oleh": "reviewed_by",
        "Didukung oleh": "supported_by",
        "Disetujui oleh": "approved_by",
        "Diterima oleh": "received_by",
    }
    style = ParagraphStyle(
        "acc",
        fontName=body_font,
        fontSize=size,
        leading=size * 1.25,
        alignment=TA_CENTER,
    )

    def cell_html(label: str) -> Paragraph:
        person = roles.get(mapping[label]) or {}
        html = (
            f"<b>{label}:</b><br/><br/>"
            f"{person.get('title') or ''}<br/><br/><br/><br/>"
            f"<u>{person.get('name') or ''}</u><br/>"
            f"{person.get('rank') or ''}"
        )
        return Paragraph(html, style)

    rows = []
    for i in range(0, len(labels), 2):
        left = cell_html(labels[i])
        right = cell_html(labels[i + 1]) if i + 1 < len(labels) else Paragraph("", style)
        rows.append([left, right])
    t = Table(rows, colWidths=[85 * mm, 85 * mm])
    t.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.75, colors.black),
                ("INNERGRID", (0, 0), (-1, -1), 0.75, colors.black),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("MINHEIGHT", (0, 0), (-1, -1), 90),
            ]
        )
    )
    return t


def _strip_placeholders(text: str) -> str:
    lines = []
    for line in (text or "").splitlines():
        if re.match(r"^\s*contoh\s*:", line, flags=re.IGNORECASE):
            continue
        lines.append(line)
    return "\n".join(lines).strip()


def _add_outline_paragraphs(document: Document, text: str, style) -> None:
    """Write body text with BI outline markers 1. / a. / i. and hanging indent."""
    for block in parse_outline_blocks(text):
        if block.get("blank"):
            document.add_paragraph("", style=style)
            continue
        if block["level"] == 0:
            p = document.add_paragraph(block["text"], style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            continue
        p = document.add_paragraph(style=style)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        level = block["level"]
        p.paragraph_format.left_indent = Cm(LEVEL_INDENT_CM[level])
        p.paragraph_format.first_line_indent = Cm(-LEVEL_HANGING_CM[level])
        run = p.add_run(f"{block['marker']} {block['text']}")
        run.font.name = "Frutiger 45 Light"


def _outline_pdf_paragraphs(text: str, body_style) -> list:
    elems = []
    for block in parse_outline_blocks(text):
        if block.get("blank"):
            elems.append(Spacer(1, 4))
            continue
        if block["level"] == 0:
            elems.append(Paragraph(block["text"].replace("&", "&amp;"), body_style))
            continue
        level = block["level"]
        left_mm = LEVEL_INDENT_CM[level] * 10
        hang_mm = LEVEL_HANGING_CM[level] * 10
        style = ParagraphStyle(
            f"outline{level}",
            parent=body_style,
            leftIndent=left_mm * mm,
            firstLineIndent=-hang_mm * mm,
            alignment=TA_JUSTIFY,
        )
        safe = f"{block['marker']} {block['text']}".replace("&", "&amp;").replace("<", "&lt;")
        elems.append(Paragraph(safe, style))
    return elems


def compose_structured_fields(sec_def: dict[str, Any], raw: dict[str, Any]) -> tuple[str, list[tuple[str, str]]]:
    """Return (legacy text, labeled rows) for structured field sections."""
    values = raw.get("fields") or {}
    rows: list[tuple[str, str]] = []
    for f in sec_def.get("fields") or []:
        val = (values.get(f["key"]) or "").strip()
        if val:
            rows.append((f["label"], val))
    text = "\n".join(f"{label}: {val}" for label, val in rows)
    return text, rows


def _margins_for(pack: dict[str, Any], doc_type_code: str) -> dict[str, float]:
    by_family = pack.get("styles", {}).get("margin_mm_by_family") or {}
    if doc_type_code in by_family:
        return by_family[doc_type_code]
    if doc_type_code.startswith("M.01") or doc_type_code == "M.01":
        return by_family.get("M.01") or pack["styles"]["margin_mm"]
    if doc_type_code.startswith("M.02") or doc_type_code == "M.02":
        return by_family.get("M.02") or pack["styles"]["margin_mm"]
    return by_family.get("MEETING_REQUEST") or pack["styles"]["margin_mm"]


def _logo_path(pack: dict[str, Any]) -> Path | None:
    rel = (pack.get("styles", {}).get("logo") or {}).get("asset")
    if not rel:
        return None
    path = ROOT / rel
    return path if path.exists() else None


def _content_model(doc: dict[str, Any]) -> dict[str, Any]:
    version = doc.get("template_version") or "1.1.0"
    tmpl = get_template(doc["type"], version)
    sections_map = {s["key"]: s for s in doc.get("sections") or []}
    rendered_sections = []
    for sec in tmpl["sections"]:
        raw = sections_map.get(sec["key"]) or {}
        if raw.get("not_needed"):
            continue
        field_rows: list[tuple[str, str]] = []
        if sec.get("input_mode") == "structured_fields":
            content, field_rows = compose_structured_fields(sec, raw)
        else:
            content = raw.get("content") or ""
            if raw.get("body_mode") != "prose" and raw.get("points"):
                numbered = points_to_content(raw["points"])
                if numbered:
                    content = numbered
            content = _strip_placeholders(content)
        table = raw.get("table")
        rendered_sections.append(
            {
                "key": sec["key"],
                "title": section_heading(sec["title"], sec.get("outline_number")),
                "outline_number": sec.get("outline_number"),
                "content": content,
                "table": table,
                "field_rows": field_rows,
                "input_mode": sec.get("input_mode"),
            }
        )
    attachments_list = [
        a for a in (doc.get("attachments_list") or [])
        if attachment_is_meaningful(a)
    ]
    meta = dict(doc.get("metadata") or {})
    if attachments_list and (not meta.get("attachments") or meta.get("attachments") in {"", "-"}):
        meta["attachments"] = format_attachment_count(len(attachments_list))
    return {
        "meta": meta,
        "type": doc["type"],
        "label": tmpl.get("label"),
        "doc_type_code": tmpl.get("doc_type_code"),
        "layout_variant": tmpl.get("layout_variant"),
        "metadata_rows": tmpl.get("metadata_rows") or [],
        "sections": rendered_sections,
        "attachments_list": attachments_list,
        "doc_id": doc.get("id"),
        "accountability": doc.get("accountability") or {},
        "signatory": doc.get("signatory") or {},
        "accountability_labels": tmpl.get("accountability", {}).get("labels") or [],
        "accountability_mode": tmpl.get("accountability", {}).get("mode"),
        "template_version": doc.get("template_version"),
        "rule_version": doc.get("rule_version"),
        "status": doc.get("status", "draft"),
    }


def build_email_body(doc: dict[str, Any]) -> str:
    model = _content_model(doc)
    parts = [f"Perihal: {model['meta'].get('subject', '')}", ""]
    for sec in model["sections"]:
        parts.append(sec["title"])
        if sec["content"]:
            parts.append(sec["content"])
        if sec.get("table") and sec["table"].get("rows"):
            cols = sec["table"].get("columns") or []
            parts.append(" | ".join(cols))
            for row in sec["table"]["rows"]:
                parts.append(" | ".join(str(row.get(c, "")) for c in cols))
        parts.append("")
    parts.append("Hormat kami,")
    parts.append((doc.get("signatory") or {}).get("name", ""))
    return "\n".join(parts).strip() + "\n"


def _clear_table_borders(table) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    # replace existing borders if any
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(borders)


def _add_kop_row(document: Document, pack: dict[str, Any], badge_code: str) -> None:
    """Logo kiri + badge kanan (satu baris) — sama seperti PDF M.02 contoh."""
    logo = _logo_path(pack)
    width_mm = float((pack.get("styles", {}).get("logo") or {}).get("width_mm") or 59.1)
    table = document.add_table(rows=1, cols=2)
    _clear_table_borders(table)
    left, right = table.rows[0].cells
    if logo:
        p = left.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run()
        run.add_picture(str(logo), width=Mm(width_mm))
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = rp.add_run(badge_code)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = body_font_family_name()
    # Kotak tipis di sekitar badge (seperti PDF)
    pPr = rp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "000000")
        pBdr.append(el)
    pPr.append(pBdr)


def _add_meta_colon_row(document: Document, label: str, value: str, style, *, bold_value: bool = False, italic_value: bool = False) -> None:
    p = document.add_paragraph(style=style)
    r1 = p.add_run(f"{label}")
    r1.bold = True if label.upper() == "PERIHAL" else False
    p.add_run(" : ")
    r2 = p.add_run(value or "")
    r2.bold = bold_value
    r2.italic = italic_value


def _add_horizontal_rule_para(document: Document, style) -> None:
    p = document.add_paragraph(style=style)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)
    p.paragraph_format.space_after = Pt(8)


def _add_logo(document: Document, pack: dict[str, Any]) -> None:
    logo = _logo_path(pack)
    if not logo:
        return
    width_mm = float((pack.get("styles", {}).get("logo") or {}).get("width_mm") or 59.1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run()
    run.add_picture(str(logo), width=Mm(width_mm))


def export_docx_bytes(doc: dict[str, Any]) -> tuple[bytes, str, str]:
    model = _content_model(doc)
    pack = load_templates(doc.get("template_version") or "1.1.0")
    margins = _margins_for(pack, model["doc_type_code"])
    sizes = pack["styles"]["font_sizes_pt"]

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(margins["top"])
    section.bottom_margin = Mm(margins["bottom"])
    section.left_margin = Mm(margins["left"])
    section.right_margin = Mm(margins["right"])

    styles = document.styles
    body_font = body_font_family_name()
    heading_font = heading_font_family_name()

    def ensure_style(name: str, base: str, size_pt: float, font_name: str) -> None:
        try:
            style = styles[name]
        except KeyError:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = styles[base]
        style.font.size = Pt(size_pt)
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)

    ensure_style("MemoTitle", "Normal", sizes["MemoTitle"], body_font)
    ensure_style("Heading1BI", "Heading 1", sizes["Heading1"], body_font)
    ensure_style("BodyBI", "Normal", sizes["Body"], body_font)
    ensure_style("MetadataBI", "Normal", sizes["Metadata"], body_font)
    ensure_style("SignatureBlock", "Normal", sizes["SignatureBlock"], body_font)

    # Best practice: logo kiri + badge kanan; No/Lamp di kiri di bawah logo
    _add_kop_row(document, pack, model["doc_type_code"] if model["doc_type_code"] != "MR" else "MR")

    meta = model["meta"]
    is_m02 = model["doc_type_code"] == "M.02" or model["layout_variant"] == "m02_satker"
    is_m01 = model["doc_type_code"] == "M.01" or model.get("layout_variant") == "m01_correspondence"
    lamp_label = "Lampiran" if is_m02 else "Lamp."

    no_p = document.add_paragraph(
        f"No. {meta.get('document_number') or '…………'}",
        style=styles["MetadataBI"],
    )
    no_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lamp_p = document.add_paragraph(
        f"{lamp_label} : {meta.get('attachments') or '-'}" if is_m02 else f"{lamp_label} {meta.get('attachments') or '-'}",
        style=styles["MetadataBI"],
    )
    lamp_p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = document.add_paragraph("MEMORANDUM" if model["doc_type_code"] != "MR" else "MEETING REQUEST")
    title.style = styles["MemoTitle"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.bold = True

    if is_m02:
        # PERIHAL (kapital) + garis → Kepada → Melalui
        subj = (meta.get("subject") or "").upper()
        _add_meta_colon_row(
            document, "PERIHAL", subj, styles["BodyBI"], bold_value=True, italic_value=True
        )
        # Garis bawah baris perihal (seperti PDF contoh)
        last = document.paragraphs[-1]
        for run in last.runs:
            run.underline = True
        _add_horizontal_rule_para(document, styles["BodyBI"])
        _add_meta_colon_row(document, "Kepada", meta.get("recipient") or "", styles["BodyBI"])
        if meta.get("via") or meta.get("melalui"):
            _add_meta_colon_row(
                document,
                "Melalui",
                meta.get("via") or meta.get("melalui") or "",
                styles["BodyBI"],
            )
    elif is_m01:
        # Kepada → Dari → Perihal → garis
        _add_meta_colon_row(document, "Kepada", meta.get("recipient") or "", styles["BodyBI"])
        _add_meta_colon_row(
            document, "Dari", meta.get("dari") or meta.get("satker") or "", styles["BodyBI"]
        )
        _add_meta_colon_row(document, "Perihal", meta.get("subject") or "", styles["BodyBI"])
        _add_horizontal_rule_para(document, styles["BodyBI"])
    else:
        document.add_paragraph(f"Hal: {meta.get('subject') or ''}", style=styles["BodyBI"])
        document.add_paragraph(f"Yth.: {meta.get('recipient') or ''}", style=styles["BodyBI"])
        _add_horizontal_rule_para(document, styles["BodyBI"])

    for idx, sec in enumerate(model["sections"], start=1):
        show_heading = not is_m01 and model["doc_type_code"] != "M.01"
        if show_heading:
            heading_text = section_heading(sec["title"], sec.get("outline_number"))
            if not sec.get("outline_number"):
                heading_text = f"{idx}. {sec['title']}"
            h = document.add_paragraph(heading_text.upper() if is_m02 else heading_text)
            h.style = styles["Heading1BI"]
            for run in h.runs:
                run.bold = True
        if sec.get("field_rows"):
            for label, val in sec["field_rows"]:
                document.add_paragraph(f"{label} : {val}", style=styles["BodyBI"])
        elif sec["content"]:
            _add_outline_paragraphs(document, sec["content"], styles["BodyBI"])
        table = sec.get("table")
        if table and table.get("rows"):
            cols = table.get("columns") or []
            t = document.add_table(rows=1, cols=len(cols))
            hdr_cells = t.rows[0].cells
            for i, c in enumerate(cols):
                hdr_cells[i].text = c
            for row in table["rows"]:
                cells = t.add_row().cells
                for i, c in enumerate(cols):
                    cells[i].text = str(row.get(c, ""))

    document.add_paragraph("")
    city = meta.get("city_date") or ""

    if model["accountability_mode"] == "grid":
        # Tanggal di atas grid (kanan), seperti PDF contoh M.02
        if city:
            dp = document.add_paragraph(city, style=styles["SignatureBlock"])
            dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_accountability_table(
            document,
            model["accountability_labels"],
            model["accountability"],
        )
    elif model["accountability_mode"] == "signatory_only" or is_m01:
        # M.01: tanggal → jabatan → ruang TTD → nama (tanpa baris unit)
        sig = model.get("signatory") or {}
        p = document.add_paragraph(city, style=styles["SignatureBlock"])
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for line, underline, bold in (
            (sig.get("title", ""), False, False),
            ("", False, False),
            ("", False, False),
            ("", False, False),
            (sig.get("name", ""), True, True),
            (sig.get("rank", ""), False, False),
        ):
            sp = document.add_paragraph(style=styles["SignatureBlock"])
            sp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = sp.add_run(line)
            run.underline = underline
            run.bold = bold
            run.font.name = body_font
            run.font.size = Pt(sizes["SignatureBlock"])
            run._element.rPr.rFonts.set(qn("w:eastAsia"), body_font)

    attachments_list = model.get("attachments_list") or []
    if attachments_list:
        document.add_paragraph("")
        h = document.add_paragraph("Lampiran", style=styles["Heading1BI"])
        for run in h.runs:
            run.bold = True
        doc_id = model.get("doc_id") or ""
        for i, item in enumerate(attachments_list, 1):
            title_a = (item.get("title") or "").strip() or f"Lampiran {i}"
            desc = (item.get("description") or "").strip()
            kind = item.get("type") or "note"
            line = f"{i}. {title_a}" + (f" — {desc}" if desc else "")
            document.add_paragraph(line, style=styles["BodyBI"])
            if kind == "table":
                table = item.get("table") or {}
                cols = table.get("columns") or []
                rows = table.get("rows") or []
                if cols and rows:
                    t = document.add_table(rows=1, cols=len(cols))
                    for ci, c in enumerate(cols):
                        t.rows[0].cells[ci].text = str(c)
                    for row in rows:
                        cells = t.add_row().cells
                        for ci, c in enumerate(cols):
                            cells[ci].text = str(row.get(c, ""))
            elif kind == "image":
                img = item.get("image") or {}
                stored = img.get("stored_name")
                if stored and doc_id:
                    try:
                        path = resolve_attachment_path(doc_id, stored)
                        document.add_picture(str(path), width=Cm(14))
                    except Exception:
                        document.add_paragraph(
                            f"[Gambar: {img.get('original_name') or stored}]",
                            style=styles["BodyBI"],
                        )

    tembusan = meta.get("tembusan") or []
    if tembusan:
        document.add_paragraph("")
        document.add_paragraph("Tembusan:", style=styles["BodyBI"])
        for i, item in enumerate(tembusan, 1):
            document.add_paragraph(f"{i}. {item}", style=styles["BodyBI"])

    apply_run_fonts(document, body_font, heading_font)

    buffer = io.BytesIO()
    document.save(buffer)
    raw = buffer.getvalue()
    checksum = hashlib.sha256(raw).hexdigest()
    filename = generate_filename(
        meta.get("satker", "SATKER"),
        meta.get("program_strategis", "PS00"),
        model["doc_type_code"],
        meta.get("subject", "Dokumen"),
        datetime.now(),
    ) + ".docx"
    return raw, filename, checksum


def export_pdf_bytes(doc: dict[str, Any]) -> tuple[bytes, str, str]:
    model = _content_model(doc)
    pack = load_templates(doc.get("template_version") or "1.1.0")
    margins = _margins_for(pack, model["doc_type_code"])
    sizes = pack["styles"]["font_sizes_pt"]
    font_paths = resolve_font_paths()

    heading_font = "Helvetica-Bold"
    body_font = "Helvetica"
    if font_paths["heading"]:
        try:
            pdfmetrics.registerFont(TTFont("BI-Heading", str(font_paths["heading"])))
            heading_font = "BI-Heading"
        except Exception:
            heading_font = "Helvetica-Bold"
    if font_paths["body"]:
        try:
            pdfmetrics.registerFont(TTFont("BI-Body", str(font_paths["body"])))
            body_font = "BI-Body"
        except Exception:
            body_font = "Helvetica"

    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margins["left"] * mm,
        rightMargin=margins["right"] * mm,
        topMargin=margins["top"] * mm,
        bottomMargin=margins["bottom"] * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "MemoTitle",
        parent=styles["Heading1"],
        fontName=body_font,
        fontSize=sizes["MemoTitle"],
        alignment=1,
        spaceAfter=12,
        spaceBefore=8,
    )
    h_style = ParagraphStyle(
        "Heading1BI",
        parent=styles["Heading2"],
        fontName=body_font,
        fontSize=sizes["Heading1"],
        spaceBefore=10,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "BodyBI",
        parent=styles["Normal"],
        fontName=body_font,
        fontSize=sizes["Body"],
        leading=sizes["Body"] * pack["styles"]["line_spacing"],
        alignment=TA_JUSTIFY,
    )
    meta_style = ParagraphStyle("MetaBI", parent=body_style, alignment=0)

    story = []
    meta = model["meta"]
    is_m02 = model["doc_type_code"] == "M.02" or model["layout_variant"] == "m02_satker"
    is_m01 = model["doc_type_code"] == "M.01" or model.get("layout_variant") == "m01_correspondence"
    badge = model["doc_type_code"] if model["doc_type_code"] != "MR" else "MR"

    logo = _logo_path(pack)
    logo_flow = ""
    if logo:
        width = float((pack.get("styles", {}).get("logo") or {}).get("width_mm") or 59.1)
        logo_flow = RLImage(str(logo), width=width * mm, height=(width * 70 / 366) * mm)
    badge_p = Paragraph(f"<b>{badge}</b>", ParagraphStyle("badge", parent=meta_style, alignment=2, borderWidth=1, borderPadding=3))
    if logo_flow:
        kop = Table([[logo_flow, badge_p]], colWidths=[120 * mm, 50 * mm])
    else:
        kop = Table([["", badge_p]], colWidths=[120 * mm, 50 * mm])
    kop.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                ("BOX", (1, 0), (1, 0), 0.75, colors.black),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    story.append(kop)
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"No. {meta.get('document_number') or '[akan diisi]'}", meta_style))
    if is_m02:
        story.append(Paragraph(f"Lampiran : {meta.get('attachments') or '-'}", meta_style))
    else:
        story.append(Paragraph(f"Lamp. {meta.get('attachments') or '-'}", meta_style))
    story.append(Paragraph("MEMORANDUM" if model["doc_type_code"] != "MR" else "MEETING REQUEST", title_style))

    if is_m02:
        subj = (meta.get("subject") or "").upper().replace("&", "&amp;")
        story.append(Paragraph(f"<b><u>PERIHAL: {subj}</u></b>", body_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(f"Kepada : {(meta.get('recipient') or '').replace('&', '&amp;')}", meta_style))
        via = meta.get("via") or meta.get("melalui") or ""
        if via:
            story.append(Paragraph(f"Melalui : {via.replace('&', '&amp;')}", meta_style))
        story.append(Spacer(1, 8))
    elif is_m01:
        data = [
            ["Kepada", ":", meta.get("recipient") or ""],
            ["Dari", ":", meta.get("dari") or meta.get("satker") or ""],
            ["Perihal", ":", meta.get("subject") or ""],
        ]
        t = Table(data, colWidths=[25 * mm, 5 * mm, 130 * mm])
        t.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), body_font),
                    ("FONTSIZE", (0, 0), (-1, -1), sizes["Body"]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 8))
    else:
        story.append(Paragraph(f"Hal: {meta.get('subject') or ''}", body_style))
        story.append(Paragraph(f"Yth.: {meta.get('recipient') or ''}", body_style))
        story.append(Spacer(1, 8))

    for sec in model["sections"]:
        if not is_m01:
            title = (sec.get("title") or "").upper() if is_m02 else (sec.get("title") or "")
            story.append(Paragraph(f"<b>{title.replace('&', '&amp;')}</b>", h_style))
        if sec.get("field_rows"):
            data = [[label, f": {val}"] for label, val in sec["field_rows"]]
            t = Table(data, colWidths=[35 * mm, 120 * mm], hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), body_font),
                        ("FONTSIZE", (0, 0), (-1, -1), sizes["Body"]),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 1),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                    ]
                )
            )
            story.append(t)
            story.append(Spacer(1, 6))
        elif sec["content"]:
            story.extend(_outline_pdf_paragraphs(sec["content"], body_style))
        table = sec.get("table")
        if table and table.get("rows"):
            cols = table.get("columns") or []
            data = [cols] + [[str(r.get(c, "")) for c in cols] for r in table["rows"]]
            t = Table(data, hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, 0), body_font),
                        ("FONTNAME", (0, 1), (-1, -1), body_font),
                        ("FONTSIZE", (0, 0), (-1, -1), sizes["Body"]),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.82, 0.82, 0.82)),
                        ("BOX", (0, 0), (-1, -1), 0.5, colors.black),
                        ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.black),
                        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.Color(0.55, 0.55, 0.55)),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(t)

    story.append(Spacer(1, 12))
    city = meta.get("city_date") or ""
    right = ParagraphStyle("sig", parent=body_style, alignment=2)

    if model["accountability_mode"] == "grid":
        if city:
            story.append(Paragraph(city.replace("&", "&amp;"), right))
            story.append(Spacer(1, 4))
        story.append(
            _accountability_pdf_table(
                model["accountability_labels"],
                model["accountability"],
                body_font,
                sizes["Body"],
            )
        )
    elif model["accountability_mode"] == "signatory_only" or is_m01:
        sig = model.get("signatory") or {}
        story.append(Paragraph((city or "").replace("&", "&amp;"), right))
        story.append(Paragraph((sig.get("title") or "").replace("&", "&amp;"), right))
        story.append(Spacer(1, 28))
        story.append(Paragraph(f"<u><b>{(sig.get('name') or '').replace('&', '&amp;')}</b></u>", right))
        if sig.get("rank"):
            story.append(Paragraph(sig.get("rank", "").replace("&", "&amp;"), right))

    attachments_list = model.get("attachments_list") or []
    if attachments_list:
        story.append(Spacer(1, 10))
        story.append(Paragraph("<b>Lampiran</b>", h_style))
        doc_id = model.get("doc_id") or ""
        for i, item in enumerate(attachments_list, 1):
            title = (item.get("title") or "").strip() or f"Lampiran {i}"
            desc = (item.get("description") or "").strip()
            kind = item.get("type") or "note"
            line = f"{i}. {title}" + (f" — {desc}" if desc else "")
            story.append(Paragraph(line.replace("&", "&amp;"), body_style))
            if kind == "table":
                table = item.get("table") or {}
                cols = table.get("columns") or []
                rows = table.get("rows") or []
                if cols and rows:
                    data = [cols] + [[str(r.get(c, "")) for c in cols] for r in rows]
                    t = Table(data, hAlign="LEFT")
                    t.setStyle(
                        TableStyle(
                            [
                                ("FONTNAME", (0, 0), (-1, 0), body_font),
                                ("FONTNAME", (0, 1), (-1, -1), body_font),
                                ("FONTSIZE", (0, 0), (-1, -1), sizes["Body"]),
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.82, 0.82, 0.82)),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ]
                        )
                    )
                    story.append(t)
                    story.append(Spacer(1, 6))
            elif kind == "image":
                img = item.get("image") or {}
                stored = img.get("stored_name")
                if stored and doc_id:
                    try:
                        path = resolve_attachment_path(doc_id, stored)
                        img_flow = RLImage(str(path))
                        img_flow._restrictSize(140 * mm, 180 * mm)
                        story.append(img_flow)
                        story.append(Spacer(1, 6))
                    except Exception:
                        story.append(
                            Paragraph(
                                f"[Gambar: {(img.get('original_name') or stored)}]",
                                body_style,
                            )
                        )

    pdf.build(story)
    raw = buffer.getvalue()
    checksum = hashlib.sha256(raw).hexdigest()
    filename = generate_filename(
        meta.get("satker", "SATKER"),
        meta.get("program_strategis", "PS00"),
        model["doc_type_code"],
        meta.get("subject", "Dokumen"),
        datetime.now(),
    ) + ".pdf"
    return raw, filename, checksum
